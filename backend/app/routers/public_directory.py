from __future__ import annotations



from io import BytesIO

import hashlib
import json
import math
import re

import threading
import time
from urllib.parse import urlencode, quote

from urllib.request import Request as UrlRequest, urlopen

from urllib.error import URLError, HTTPError



from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse

from pydantic import BaseModel

from sqlalchemy.orm import Session



from app.db.session import get_db
from app.auth import (
    get_current_user,
    log_workspace_activity,
    require_company_write_access,
)

from app.db_models.entities import (

    Company,

    SupplyNode,

    SupplyRoute,
    User,

)





def require_company_access(
    company_id: int,
    request: Request,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_company_write_access(company_id, current_user, db)
    tracked_changes = {
        "import-public-data": "Imported public company network data.",
        "personalize-demo-network": "Updated the synthetic demo network profile.",
        "create-demo-network": "Created a synthetic demo supply-chain network.",
        "import-excel": "Imported a supply-chain network from an Excel workbook.",
    }
    path_parts = request.url.path.rstrip("/").split("/")
    action = path_parts[-2] if path_parts[-1].isdigit() else path_parts[-1]
    if action in tracked_changes:
        log_workspace_activity(
            db,
            company_id,
            current_user,
            "network_import",
            tracked_changes[action],
        )


router = APIRouter(

    prefix="/public-directory",

    tags=["Public Company Directory"],

)

COUNTRY_ENTITY_IDS = {
    "czechia": "Q213",
    "hong kong": "Q8646",
    "taiwan": "Q865",
    "united states": "Q30",
}
COUNTRY_FEATURED_COMPANIES = {
    "argentina": (("YPF", "Energy"), ("Mercado Libre", "E-commerce and technology"), ("Arcor", "Food manufacturing")),
    "australia": (("BHP", "Mining"), ("Qantas", "Airline"), ("Wesfarmers", "Retail and industrials")),
    "austria": (("OMV", "Energy"), ("Erste Group", "Banking"), ("Red Bull", "Beverages")),
    "bangladesh": (("Grameenphone", "Telecommunications"), ("Beximco", "Manufacturing and pharmaceuticals"), ("Square Pharmaceuticals", "Pharmaceuticals")),
    "belgium": (("AB InBev", "Beverages"), ("Solvay", "Chemicals"), ("Umicore", "Materials and recycling")),
    "brazil": (("Petrobras", "Energy"), ("Vale", "Mining"), ("Embraer", "Aerospace")),
    "canada": (("Shopify", "E-commerce software"), ("Bombardier", "Aerospace"), ("Canadian National Railway", "Rail freight")),
    "chile": (("Codelco", "Mining"), ("Falabella", "Retail"), ("LATAM Airlines Group", "Airline")),
    "china": (("Alibaba Group", "E-commerce and technology"), ("BYD", "Automotive and batteries"), ("COSCO Shipping", "Shipping and logistics"), ("Tencent", "Technology"), ("Lenovo", "Electronics")),
    "colombia": (("Ecopetrol", "Energy"), ("Grupo Aval", "Banking"), ("Bancolombia", "Banking")),
    "czechia": (("Škoda Auto", "Automotive"), ("ČEZ Group", "Energy"), ("PPF Group", "Investment and financial services")),
    "denmark": (("Maersk", "Shipping and logistics"), ("Novo Nordisk", "Pharmaceuticals"), ("LEGO Group", "Toys and manufacturing")),
    "egypt": (("Elsewedy Electric", "Electrical equipment"), ("Orascom Construction", "Construction"), ("Commercial International Bank", "Banking")),
    "finland": (("Nokia", "Telecommunications"), ("KONE", "Elevators and escalators"), ("Wärtsilä", "Marine and energy technology")),
    "france": (("Renault", "Automotive"), ("L'Oréal", "Consumer goods"), ("Schneider Electric", "Energy management"), ("Sanofi", "Pharmaceuticals"), ("TotalEnergies", "Energy")),
    "germany": (("Siemens", "Industrial technology"), ("BMW", "Automotive"), ("SAP", "Enterprise software"), ("Volkswagen Group", "Automotive"), ("Deutsche Telekom", "Telecommunications")),
    "greece": (("Aegean Airlines", "Airline"), ("HELLENiQ Energy", "Energy"), ("OTE Group", "Telecommunications")),
    "hong kong": (("Cathay Pacific", "Airline"), ("CLP Holdings", "Energy"), ("MTR Corporation", "Rail transport")),
    "hungary": (("MOL Group", "Energy"), ("OTP Bank", "Banking"), ("Richter Gedeon", "Pharmaceuticals")),
    "india": (("Tata Motors", "Automotive"), ("Reliance Industries", "Energy and retail"), ("Mahindra Group", "Automotive and manufacturing"), ("Infosys", "Information technology"), ("Tata Consultancy Services", "Information technology")),
    "indonesia": (("Pertamina", "Energy"), ("Astra International", "Automotive and services"), ("GoTo Group", "Technology and delivery")),
    "ireland": (("Ryanair", "Airline"), ("Kerry Group", "Food and ingredients"), ("CRH", "Building materials")),
    "israel": (("Teva Pharmaceutical Industries", "Pharmaceuticals"), ("Check Point Software", "Cybersecurity"), ("Elbit Systems", "Aerospace and defense")),
    "italy": (("Ferrari", "Automotive"), ("Enel", "Energy"), ("Eni", "Energy")),
    "japan": (("Toyota", "Automotive"), ("Sony Group", "Electronics and entertainment"), ("Mitsubishi Heavy Industries", "Industrial manufacturing"), ("Honda", "Automotive"), ("Panasonic", "Electronics")),
    "malaysia": (("Petronas", "Energy"), ("Maybank", "Banking"), ("Top Glove", "Medical supplies")),
    "mexico": (("Grupo Bimbo", "Food manufacturing"), ("América Móvil", "Telecommunications"), ("Cemex", "Building materials")),
    "netherlands": (("ASML", "Semiconductor equipment"), ("Philips", "Health technology"), ("Heineken", "Beverages")),
    "new zealand": (("Fonterra", "Dairy and food"), ("Fisher & Paykel Healthcare", "Medical equipment"), ("Air New Zealand", "Airline")),
    "nigeria": (("Dangote Cement", "Building materials"), ("Nigerian Breweries", "Beverages"), ("MTN Nigeria", "Telecommunications")),
    "norway": (("Equinor", "Energy"), ("Yara International", "Fertilizers"), ("Norsk Hydro", "Aluminium and energy")),
    "pakistan": (("Engro Corporation", "Energy and chemicals"), ("Lucky Cement", "Building materials"), ("Habib Bank", "Banking")),
    "philippines": (("Jollibee Foods", "Food service"), ("SM Investments", "Retail and property"), ("San Miguel Corporation", "Food and manufacturing")),
    "poland": (("Orlen", "Energy"), ("CD Projekt", "Software and games"), ("Allegro", "E-commerce")),
    "portugal": (("EDP", "Energy"), ("Galp", "Energy"), ("Jerónimo Martins", "Food retail")),
    "qatar": (("Qatar Airways", "Airline"), ("Ooredoo", "Telecommunications"), ("QatarEnergy", "Energy")),
    "romania": (("Dacia", "Automotive"), ("OMV Petrom", "Energy"), ("Banca Transilvania", "Banking")),
    "saudi arabia": (("Saudi Aramco", "Energy"), ("SABIC", "Chemicals"), ("stc Group", "Telecommunications")),
    "singapore": (("DBS Bank", "Banking"), ("Singapore Airlines", "Airline"), ("PSA International", "Port operations and logistics"), ("Singtel", "Telecommunications"), ("Grab", "Technology and mobility")),
    "south africa": (("Naspers", "Technology and media"), ("Shoprite", "Retail"), ("Sasol", "Energy and chemicals")),
    "south korea": (("Samsung Electronics", "Electronics"), ("Hyundai Motor", "Automotive"), ("LG Electronics", "Electronics")),
    "spain": (("Inditex", "Apparel retail"), ("Iberdrola", "Energy"), ("Banco Santander", "Banking")),
    "sweden": (("Volvo Group", "Automotive and transport"), ("Ericsson", "Telecommunications"), ("H&M", "Apparel retail")),
    "switzerland": (("Nestlé", "Food and beverages"), ("Roche", "Pharmaceuticals"), ("UBS", "Banking")),
    "taiwan": (("TSMC", "Semiconductors"), ("Foxconn", "Electronics manufacturing"), ("ASUS", "Electronics")),
    "thailand": (("PTT", "Energy"), ("Charoen Pokphand Group", "Food and retail"), ("Thai Beverage", "Beverages")),
    "turkey": (("Turkish Airlines", "Airline"), ("Koç Holding", "Industrial and consumer businesses"), ("Şişecam", "Glass manufacturing")),
    "united arab emirates": (("Emirates", "Airline"), ("ADNOC", "Energy"), ("DP World", "Port operations and logistics")),
    "united kingdom": (("Rolls-Royce Holdings", "Aerospace and power systems"), ("Tesco", "Retail"), ("BT Group", "Telecommunications"), ("AstraZeneca", "Pharmaceuticals"), ("BP", "Energy")),
    "united states": (("Apple", "Technology"), ("Amazon", "E-commerce and cloud computing"), ("Boeing", "Aerospace"), ("Microsoft", "Technology"), ("Walmart", "Retail")),
    "vietnam": (("Vingroup", "Industrial and consumer businesses"), ("Viettel", "Telecommunications"), ("Vinamilk", "Food manufacturing")),
}
COMPANY_DIRECTORY_CACHE_TTL = 900
_company_directory_cache: dict[tuple[str, str], tuple[float, list[dict]]] = {}
_company_directory_cache_lock = threading.Lock()





FALLBACK_COUNTRIES = [

    "Argentina",

    "Australia",

    "Austria",

    "Bangladesh",

    "Belgium",

    "Brazil",

    "Canada",

    "Chile",

    "China",

    "Colombia",

    "Czechia",

    "Denmark",

    "Egypt",

    "Finland",

    "France",

    "Germany",

    "Greece",

    "Hong Kong",

    "Hungary",

    "India",

    "Indonesia",

    "Ireland",

    "Israel",

    "Italy",

    "Japan",

    "Malaysia",

    "Mexico",

    "Netherlands",

    "New Zealand",

    "Nigeria",

    "Norway",

    "Pakistan",

    "Philippines",

    "Poland",

    "Portugal",

    "Qatar",

    "Romania",

    "Saudi Arabia",

    "Singapore",

    "South Africa",

    "South Korea",

    "Spain",

    "Sweden",

    "Switzerland",

    "Taiwan",

    "Thailand",

    "Turkey",

    "United Arab Emirates",

    "United Kingdom",

    "United States",

    "Vietnam",

]

COUNTRY_DEMO_NETWORK_CITIES = {
    "argentina": ("Córdoba", "Rosario", "Buenos Aires", "Mar del Plata", "Mendoza", "Córdoba", "Buenos Aires"),
    "australia": ("Melbourne", "Sydney", "Melbourne", "Brisbane", "Perth", "Adelaide", "Sydney"),
    "austria": ("Graz", "Linz", "Vienna", "Linz", "Salzburg", "Graz", "Vienna"),
    "bangladesh": ("Chattogram", "Dhaka", "Chattogram", "Mongla", "Dhaka", "Khulna", "Dhaka"),
    "belgium": ("Antwerp", "Brussels", "Antwerp", "Zeebrugge", "Ghent", "Liège", "Brussels"),
    "brazil": ("Campinas", "São Paulo", "Santos", "Rio de Janeiro", "Brasília", "Curitiba", "São Paulo"),
    "canada": ("Toronto", "Montreal", "Vancouver", "Halifax", "Calgary", "Ottawa", "Toronto"),
    "chile": ("Santiago", "Concepción", "Valparaíso", "San Antonio", "Santiago", "Valparaíso", "Santiago"),
    "china": ("Shanghai", "Suzhou", "Shanghai", "Shenzhen", "Wuhan", "Chengdu", "Beijing"),
    "colombia": ("Medellín", "Bogotá", "Cartagena", "Buenaventura", "Cali", "Barranquilla", "Bogotá"),
    "czechia": ("Brno", "Ostrava", "Prague", "Pilsen", "Brno", "Ostrava", "Prague"),
    "denmark": ("Aarhus", "Odense", "Aarhus", "Copenhagen", "Aalborg", "Odense", "Copenhagen"),
    "egypt": ("Alexandria", "Cairo", "Alexandria", "Port Said", "Cairo", "Giza", "Cairo"),
    "finland": ("Tampere", "Helsinki", "Helsinki", "Turku", "Tampere", "Oulu", "Helsinki"),
    "france": ("Lyon", "Toulouse", "Marseille", "Le Havre", "Lille", "Bordeaux", "Paris"),
    "germany": ("Stuttgart", "Munich", "Hamburg", "Bremerhaven", "Frankfurt", "Leipzig", "Berlin"),
    "greece": ("Thessaloniki", "Athens", "Piraeus", "Patras", "Athens", "Larissa", "Athens"),
    "hong kong": ("Kowloon", "Kowloon", "Hong Kong", "Hong Kong", "Kowloon", "Sha Tin", "Hong Kong"),
    "hungary": ("Szeged", "Győr", "Budapest", "Győr", "Debrecen", "Szeged", "Budapest"),
    "india": ("Pune", "Pune", "Navi Mumbai", "Chennai", "Hyderabad", "Vijayawada", "Vijayawada"),
    "indonesia": ("Bandung", "Jakarta", "Jakarta", "Surabaya", "Surabaya", "Semarang", "Jakarta"),
    "ireland": ("Cork", "Galway", "Dublin", "Cork", "Limerick", "Galway", "Dublin"),
    "israel": ("Haifa", "Tel Aviv", "Haifa", "Ashdod", "Jerusalem", "Beersheba", "Tel Aviv"),
    "italy": ("Turin", "Milan", "Genoa", "Naples", "Bologna", "Verona", "Rome"),
    "japan": ("Nagoya", "Osaka", "Yokohama", "Kobe", "Tokyo", "Fukuoka", "Tokyo"),
    "malaysia": ("Penang", "Kuala Lumpur", "Port Klang", "Johor Bahru", "Kuala Lumpur", "Kuching", "Kuala Lumpur"),
    "mexico": ("Monterrey", "Mexico City", "Veracruz", "Manzanillo", "Guadalajara", "Puebla", "Mexico City"),
    "netherlands": ("Eindhoven", "Utrecht", "Rotterdam", "Amsterdam", "Tilburg", "Groningen", "Amsterdam"),
    "new zealand": ("Christchurch", "Auckland", "Auckland", "Wellington", "Hamilton", "Dunedin", "Auckland"),
    "nigeria": ("Kano", "Lagos", "Lagos", "Port Harcourt", "Abuja", "Ibadan", "Lagos"),
    "norway": ("Trondheim", "Oslo", "Bergen", "Oslo", "Oslo", "Stavanger", "Oslo"),
    "pakistan": ("Faisalabad", "Lahore", "Karachi", "Gwadar", "Lahore", "Islamabad", "Karachi"),
    "philippines": ("Cebu City", "Manila", "Manila", "Cebu City", "Davao City", "Iloilo City", "Manila"),
    "poland": ("Wrocław", "Łódź", "Gdańsk", "Szczecin", "Warsaw", "Kraków", "Warsaw"),
    "portugal": ("Porto", "Coimbra", "Lisbon", "Porto", "Lisbon", "Braga", "Lisbon"),
    "qatar": ("Al Rayyan", "Doha", "Doha", "Al Wakrah", "Doha", "Al Khor", "Doha"),
    "romania": ("Cluj-Napoca", "Bucharest", "Constanța", "Galați", "Bucharest", "Timișoara", "Bucharest"),
    "saudi arabia": ("Jeddah", "Riyadh", "Jeddah", "Dammam", "Riyadh", "Medina", "Riyadh"),
    "singapore": ("Jurong", "Tuas", "Singapore", "Singapore", "Jurong", "Woodlands", "Singapore"),
    "south africa": ("Johannesburg", "Pretoria", "Durban", "Cape Town", "Johannesburg", "Gqeberha", "Cape Town"),
    "south korea": ("Daegu", "Daejeon", "Busan", "Incheon", "Seoul", "Gwangju", "Seoul"),
    "spain": ("Barcelona", "Madrid", "Valencia", "Bilbao", "Zaragoza", "Seville", "Madrid"),
    "sweden": ("Gothenburg", "Stockholm", "Gothenburg", "Malmö", "Stockholm", "Uppsala", "Stockholm"),
    "switzerland": ("Basel", "Zürich", "Basel", "Geneva", "Bern", "Lausanne", "Zürich"),
    "taiwan": ("Taichung", "Taipei", "Kaohsiung", "Keelung", "Taichung", "Tainan", "Taipei"),
    "thailand": ("Chiang Mai", "Bangkok", "Laem Chabang", "Bangkok", "Bangkok", "Phuket", "Bangkok"),
    "turkey": ("Bursa", "Ankara", "Istanbul", "Izmir", "Ankara", "Gaziantep", "Istanbul"),
    "united arab emirates": ("Sharjah", "Abu Dhabi", "Jebel Ali", "Khalifa Port", "Dubai", "Sharjah", "Dubai"),
    "united kingdom": ("Manchester", "Birmingham", "Southampton", "Liverpool", "Leeds", "Bristol", "London"),
    "united states": ("Chicago", "Detroit", "Los Angeles", "New York", "Dallas", "Atlanta", "New York"),
    "vietnam": ("Da Nang", "Hanoi", "Hai Phong", "Ho Chi Minh City", "Ho Chi Minh City", "Can Tho", "Hanoi"),
}




COMPANY_DESCRIPTION_KEYWORDS = (

    "company",

    "corporation",

    "business",

    "manufacturer",

    "manufacturing",

    "enterprise",

    "group",

    "conglomerate",

    "bank",

    "airline",

    "automotive",

    "pharmaceutical",

    "technology",

    "logistics",

    "retail",

    "energy",

    "telecommunications",

    "industrial",

)





class PublicCompanyImport(BaseModel):

    external_id: str

    name: str

    description: str | None = None

    source_url: str | None = None

    country: str





def _fetch_json(url: str) -> dict | list:

    request = UrlRequest(

        url,

        headers={

            "User-Agent": (

                "SupplyChainTwinAI/0.6 "

                "(local educational digital twin)"

            )

        },

    )



    with urlopen(

        request,

        timeout=15,

    ) as response:

        return json.loads(

            response.read().decode(

                "utf-8"

            )

        )


def _country_company_results(country_name: str, query: str) -> list[dict]:
    cache_key = (country_name.casefold(), query.casefold())
    now = time.monotonic()

    with _company_directory_cache_lock:
        cached_result = _company_directory_cache.get(cache_key)
        if cached_result and now - cached_result[0] < COMPANY_DIRECTORY_CACHE_TTL:
            return cached_result[1]

    if not query:
        featured_results = [
            {
                "external_id": (
                    f"LOCAL-{re.sub(r'[^a-z0-9]+', '-', country_name.casefold()).strip('-')}"
                    f"-{index + 1}"
                ),
                "name": name,
                "description": f"Curated company example · {industry}",
                "source_url": None,
                "country": country_name,
            }
            for index, (name, industry) in enumerate(
                COUNTRY_FEATURED_COMPANIES.get(country_name.casefold(), ())
            )
        ]
        with _company_directory_cache_lock:
            _company_directory_cache[cache_key] = (
                time.monotonic(),
                featured_results,
            )
        return featured_results

    country_id = COUNTRY_ENTITY_IDS.get(country_name.casefold())
    country_filter = (
        f"BIND(wd:{country_id} AS ?country)"
        if country_id
        else (
            "?country wdt:P31 wd:Q6256; "
            f"rdfs:label {json.dumps(country_name)}@en."
        )
    )
    search_filter = (
        f'FILTER(CONTAINS(LCASE(STR(?companyLabel)), '
        f'LCASE({json.dumps(query)})))'
        if query
        else ""
    )
    sparql_query = f"""
        SELECT ?company ?companyLabel WHERE {{
            {country_filter}
            ?company wdt:P31 wd:Q4830453; wdt:P17 ?country;
                rdfs:label ?companyLabel.
            FILTER(LANG(?companyLabel) = "en")
            {search_filter}
        }}
        LIMIT 100
    """
    params = urlencode(
        {
            "query": sparql_query,
            "format": "json",
        }
    )
    request = UrlRequest(
        f"https://query.wikidata.org/sparql?{params}",
        headers={
            "User-Agent": (
                "SupplyChainTwinAI/0.6 "
                "(local educational digital twin)"
            ),
            "Accept": "application/sparql-results+json",
        },
    )

    try:
        with urlopen(request, timeout=20) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, OSError, json.JSONDecodeError) as error:
        raise HTTPException(
            status_code=502,
            detail=(
                "The public company directory could not be reached. "
                "Please retry in a moment."
            ),
        ) from error

    rows = []
    for binding in payload.get("results", {}).get("bindings", []):
        entity_url = binding.get("company", {}).get("value", "")
        entity_id = entity_url.rsplit("/", 1)[-1]
        name = binding.get("companyLabel", {}).get("value", "").strip()

        if (
            not re.fullmatch(r"Q\d+", entity_id)
            or not name
            or ":" in name
        ):
            continue

        rows.append(
            {
                "external_id": entity_id,
                "name": name,
                "description": None,
                "source_url": f"https://www.wikidata.org/wiki/{entity_id}",
                "country": country_name,
            }
        )

    with _company_directory_cache_lock:
        _company_directory_cache[cache_key] = (time.monotonic(), rows)

    return rows




@router.get("/countries")

def list_countries():

    try:

        data = _fetch_json(

            "https://restcountries.com/v3.1/all?fields=name"

        )



        countries = sorted(

            {

                item.get(

                    "name",

                    {},

                ).get(

                    "common"

                )

                for item in data

                if item.get(

                    "name",

                    {},

                ).get(

                    "common"

                )

            }

        )



        if countries:

            return {

                "countries": countries,

                "source": "REST Countries",

            }

    except Exception:

        pass



    return {

        "countries":

            FALLBACK_COUNTRIES,

        "source":

            "built-in fallback",

    }





COMPANY_EXCLUDE_PATTERNS = [
    r"\bfamily name\b",
    r"\bsurname\b",
    r"\bgiven name\b",
    r"\bpatronymic\b",
    r"\btown\b",
    r"\bvillage\b",
    r"\bcity\b",
    r"\bmunicipality\b",
    r"\bcommune\b",
    r"\bdistrict\b",
    r"\bhuman settlement\b",
    r"\badministrative territorial entity\b",
    r"\bcapital of\b",
    r"\bpopulated place\b",
    r"\bisland\b",
    r"\briver\b",
    r"\bmountain\b",
    r"\bcrater\b",
    r"\blake\b",
    r"\bprovince\b",
    r"\bstate of\b",
    r"\bfederal subject\b",
    r"\bcosmopolis\b",
    r"\bmetropolis\b",
    r"\bsettlement\b",
    r"\bwikimedia\b",
    r"\bwikipedia\b",
    r"\bdisambiguation\b",
    r"\bcategory\b",
    r"\btemplate\b",
    r"\bfootballer\b",
    r"\bcricketer\b",
    r"\bpolitician\b",
    r"\bactor\b",
    r"\bactress\b",
    r"\bsinger\b",
    r"\brapper\b",
    r"\bartist\b",
    r"\bwriter\b",
    r"\bking of\b",
    r"\bqueen of\b",
    r"\bemperor\b",
    r"\bhuman\b",
    r"\bperson\b",
    r"\bathlete\b",
    r"\bswimmer\b",
    r"\bchemist\b",
    r"\bphysicist\b",
    r"\bbiologist\b",
    r"\bcomposer\b",
    r"\balbum\b",
    r"\bsong\b",
    r"\bfilm\b",
    r"\bdocumentary\b",
    r"\bmovie\b",
    r"\bnovel\b",
    r"\brecord label\b",
    r"\bstreaming service\b",
    r"\bvideo streaming service\b",
    r"\bsmartwatch\b",
    r"\bsmartphone\b",
    r"\bmobile app\b",
    r"\bapp\b",
    r"\bretail store\b",
    r"\bstore chain\b",
    r"\bshopping center\b",
    r"\bwebsite\b",
    r"\bmagazine\b",
    r"\bnewspaper\b",
    r"\btelevision series\b",
    r"\bepisode\b",
    r"\bvideo game\b",
    r"\bconsole\b",
    r"\basteroid\b",
    r"\bgalaxy\b",
    r"\bcomet\b",
    r"\bplant species\b",
    r"\banimal species\b",
    r"\bspecies of\b",
    r"\bgenus of\b",
    r"\blanguage spoken\b",
    r"\bturkic language\b",
    r"\bethnic group\b",
    r"\bpeoples\b",
    r"\bresearch institute\b",
    r"\buniversity\b",
    r"\bcollege\b",
    r"\bschool\b",
    r"\bacademy\b",
    r"\bhospital\b",
]

COMPANY_PREFERRED_KEYWORDS = [
    "company",
    "corporation",
    "business",
    "manufacturer",
    "manufacturing",
    "conglomerate",
    "enterprise",
    "group",
    "bank",
    "automotive",
    "technology",
    "pharmaceutical",
    "logistics",
    "retail",
    "retailer",
    "energy",
    "industrial",
    "telecommunications",
    "airline",
    "firm",
    "operator",
    "holding",
    "producer",
    "supplier",
    "consulting",
    "commercial",
    "multinational",
    "semiconductor",
    "electronics",
    "chemicals",
    "aerospace",
    "maritime",
    "shipping",
    "freight",
]

COMPANY_NAME_CORPORATE_SUFFIXES = [
    r"\bgroup\b",
    r"\bmotors\b",
    r"\bsteel\b",
    r"\bpower\b",
    r"\btechnologies\b",
    r"\bsolutions\b",
    r"\bindustries\b",
    r"\bholdings\b",
    r"\bltd\b",
    r"\blimited\b",
    r"\binc\b",
    r"\bcorp\b",
    r"\bcorporation\b",
    r"\bbank\b",
    r"\bairlines\b",
    r"\benergy\b",
    r"\blogistics\b",
    r"\bpharmaceuticals\b",
    r"\bservices\b",
]


PRODUCT_EXCLUDE_PATTERNS = [
    r"\brecord label\b",
    r"\bstreaming service\b",
    r"\bvideo streaming service\b",
    r"\bsmartwatch\b",
    r"\bsmartphone\b",
    r"\bmobile app\b",
    r"\bretail store\b",
    r"\bstore chain\b",
    r"\bshopping center\b",
    r"\bwebsite\b",
    r"\bmagazine\b",
    r"\bnewspaper\b",
    r"\bbrand of\b",
    r"\bvideo game\b",
    r"\bmovie\b",
    r"\bfilm\b",
    r"\bseries\b",
]


def _looks_like_company(label: str, description: str) -> bool:
    if ":" in label:
        return False

    label_lower = label.lower()
    description_lower = (description or "").lower()
    combined_text = f"{label_lower} {description_lower}"

    has_company_indicator = any(
        kw in description_lower for kw in COMPANY_PREFERRED_KEYWORDS
    ) or any(
        re.search(suffix, label_lower)
        for suffix in COMPANY_NAME_CORPORATE_SUFFIXES
    )

    if any(
        re.search(pattern, combined_text)
        for pattern in COMPANY_EXCLUDE_PATTERNS
    ):
        return False

    if any(
        re.search(pattern, combined_text)
        for pattern in PRODUCT_EXCLUDE_PATTERNS
    ):
        return any(
            re.search(suffix, label_lower)
            for suffix in COMPANY_NAME_CORPORATE_SUFFIXES
        ) or any(
            kw in description_lower
            for kw in (
                "company",
                "corporation",
                "business",
                "manufacturer",
                "manufacturing",
                "conglomerate",
                "enterprise",
                "group",
                "holding",
                "bank",
                "retailer",
            )
        )

    return has_company_indicator


def _company_name_key(name: str) -> str:
    without_suffix = re.sub(
        r"\b(limited|ltd|incorporated|inc|corporation|corp|company|co)\b",
        "",
        name.casefold(),
    )
    return re.sub(r"[^\w]", "", without_suffix)


@router.get("/companies")
def search_companies(
    country: str,
    q: str = "",
    db: Session = Depends(get_db),
):
    query = q.strip()
    country_name = country.strip()

    if not country_name:
        raise HTTPException(
            status_code=400,
            detail="Select a country to browse companies.",
        )

    merged_candidates: dict[str, dict] = {}
    local_company_count = 0
    for company in db.query(Company).all():
        description = company.description or ""
        country_match = re.search(r"Country:\s*([^\n]+)", description)
        entity_match = re.search(r"Public directory ID:\s*(Q\d+)", description)
        public_description_match = re.search(
            r"Public description:\s*([^\n]+)",
            description,
        )

        if (
            not country_match
            or country_match.group(1).strip().casefold() != country_name.casefold()
            or not entity_match
        ):
            continue

        entity_id = entity_match.group(1)
        public_description = (
            public_description_match.group(1).strip()
            if public_description_match
            else ""
        )
        label = company.name.strip()
        label_lower = label.casefold()
        if not _looks_like_company(label, public_description):
            continue
        if query and query.casefold() not in label_lower:
            continue

        merged_candidates[entity_id] = {
            "score": 100,
            "external_id": entity_id,
            "name": label,
            "description": public_description or None,
            "source_url": f"https://www.wikidata.org/wiki/{entity_id}",
            "country": country_name,
        }
        local_company_count += 1

    try:
        directory_results = _country_company_results(country_name, query)
    except HTTPException:
        directory_results = [
            company
            for company in _country_company_results(country_name, "")
            if not query or query.casefold() in company["name"].casefold()
        ]
        search_note = (
            "The public company directory is temporarily unavailable. "
            "Showing curated company examples and saved companies instead."
        )
    else:
        if all(
            company["external_id"].startswith("LOCAL-")
            for company in directory_results
        ):
            search_note = (
                "Showing curated company examples for the selected country. "
                "Search by name to check the public directory for more companies."
            )
        else:
            search_note = (
                "Companies are filtered by their registered country."
            )

    if query and any(
        not company["external_id"].startswith("LOCAL-")
        for company in directory_results
    ):
        live_names = {
            _company_name_key(company["name"])
            for company in directory_results
        }
        curated_matches = [
            company
            for company in _country_company_results(country_name, "")
            if query.casefold() in company["name"].casefold()
            and _company_name_key(company["name"]) not in live_names
        ]
        directory_results.extend(curated_matches)
        if curated_matches:
            search_note += " Curated matching examples are included."

    for company in directory_results:
        if company["external_id"] in merged_candidates or any(
            _company_name_key(company["name"]) == _company_name_key(item["name"])
            for item in merged_candidates.values()
        ):
            continue

        name_lower = company["name"].casefold()
        description_lower = (company["description"] or "").casefold()
        matching_keyword_count = sum(
            1
            for keyword in COMPANY_PREFERRED_KEYWORDS
            if keyword in description_lower
        )
        score = min(matching_keyword_count * 5, 20)

        if query and name_lower == query.casefold():
            score += 40
        elif query and name_lower.startswith(query.casefold()):
            score += 25
        elif query:
            score += 15

        has_corporate_label = any(
            re.search(suffix, name_lower)
            for suffix in COMPANY_NAME_CORPORATE_SUFFIXES
        )
        if has_corporate_label:
            score += 10

        merged_candidates[company["external_id"]] = {
            **company,
            "score": score,
        }

    sorted_results = sorted(
        merged_candidates.values(),
        key=lambda item: (-item["score"], item["name"].casefold()),
    )[:100]
    clean_results = [
        {
            "external_id": item["external_id"],
            "name": item["name"],
            "description": item["description"],
            "source_url": item["source_url"],
            "country": item["country"],
        }
        for item in sorted_results
    ]

    if local_company_count and directory_results:
        search_note += (
            f" {local_company_count} saved local "
            f"{'company is' if local_company_count == 1 else 'companies are'} "
            "included."
        )

    return {
        "country": country_name,
        "query": query,
        "results": clean_results,
        "source": (
            "Wikidata"
            if any(
                not company["external_id"].startswith("LOCAL-")
                for company in directory_results
            )
            else "Curated examples"
        ),
        "note": search_note,
    }


@router.post(
    "/import-company",
    dependencies=[Depends(get_current_user)],
)

def import_company(

    payload: PublicCompanyImport,

    current_user: User = Depends(get_current_user),

    db: Session = Depends(

        get_db

    ),

):

    existing = (

        db.query(Company)

        .filter(

            Company.name

            == payload.name

        )

        .first()

    )



    if existing is not None:
        if existing.owner_id != current_user.id:
            raise HTTPException(
                status_code=409,
                detail="A workspace with this company name already exists in another account.",
            )
        return {

            "created": False,

            "company": {

                "id":

                    existing.id,

                "name":

                    existing.name,

                "description":

                    existing.description,

                "is_demo":

                    existing.is_demo,

            },

        }



    description_lines = [

        f"Country: {payload.country}",

        (

            "Public directory ID: "

            f"{payload.external_id}"

        ),

    ]



    if payload.source_url:

        description_lines.append(

            "Public source: "

            f"{payload.source_url}"

        )



    if payload.description:

        description_lines.append(

            "Public description: "

            f"{payload.description}"

        )



    company = Company(

        name=payload.name,

        description="\n".join(

            description_lines

        ),

        is_demo=False,
        owner_id=current_user.id,
    )



    db.add(company)

    db.commit()

    db.refresh(company)



    return {

        "created": True,

        "company": {

            "id": company.id,

            "name": company.name,

            "description":

                company.description,

            "is_demo":

                company.is_demo,

        },

    }





def _extract_public_id(

    description: str | None,

) -> str | None:

    if not description:

        return None



    match = re.search(

        r"Public directory ID:\s*(Q\d+)",

        description,

    )



    return (

        match.group(1)

        if match

        else None

    )





def _claim_value(

    claims: dict,

    property_id: str,

):

    values = []



    for claim in claims.get(

        property_id,

        [],

    ):

        mainsnak = claim.get(

            "mainsnak",

            {},

        )



        datavalue = mainsnak.get(

            "datavalue",

            {},

        )



        value = datavalue.get(

            "value"

        )



        if value is not None:

            values.append(value)



    return values





def _resolve_entity_labels(

    ids: list[str],

) -> dict[str, str]:

    clean = [

        item

        for item in ids

        if isinstance(

            item,

            str,

        )

        and item.startswith("Q")

    ]



    if not clean:

        return {}



    params = urlencode(

        {

            "action":

                "wbgetentities",

            "ids":

                "|".join(

                    clean[:50]

                ),

            "props":

                "labels",

            "languages":

                "en",

            "format":

                "json",

        }

    )



    data = _fetch_json(

        "https://www.wikidata.org/"

        f"w/api.php?{params}"

    )



    labels = {}



    for entity_id, entity in (

        data.get(

            "entities",

            {}

        ).items()

    ):

        label = (

            entity.get(

                "labels",

                {},

            )

            .get(

                "en",

                {},

            )

            .get(

                "value"

            )

        )



        if label:

            labels[

                entity_id

            ] = label



    return labels





@router.post(

    "/import-public-data/{company_id}",
    dependencies=[Depends(require_company_access)],

)

def import_public_data(

    company_id: int,

    db: Session = Depends(

        get_db

    ),

):

    company = db.get(

        Company,

        company_id,

    )



    if company is None:

        raise HTTPException(

            status_code=404,

            detail=(

                "Company workspace "

                "not found."

            ),

        )



    entity_id = (

        _extract_public_id(

            company.description

        )

    )



    if not entity_id:

        raise HTTPException(

            status_code=400,

            detail=(

                "This workspace was not "

                "created from a public "

                "directory result. Search "

                "and select the company "

                "first."

            ),

        )



    try:

        entity_data = _fetch_json(

            "https://www.wikidata.org/"

            f"wiki/Special:EntityData/"

            f"{entity_id}.json"

        )

    except Exception as error:

        raise HTTPException(

            status_code=502,

            detail=(

                "Unable to download "

                "public company data: "

                f"{error}"

            ),

        )



    entity = (

        entity_data.get(

            "entities",

            {}

        ).get(

            entity_id,

            {}

        )

    )



    claims = entity.get(

        "claims",

        {}

    )



    websites = _claim_value(

        claims,

        "P856",

    )



    inception_values = _claim_value(

        claims,

        "P571",

    )



    hq_values = _claim_value(

        claims,

        "P159",

    )



    industry_values = _claim_value(

        claims,

        "P452",

    )



    hq_ids = [

        value.get("id")

        for value in hq_values

        if isinstance(

            value,

            dict,

        )

        and value.get("id")

    ]



    industry_ids = [

        value.get("id")

        for value

        in industry_values

        if isinstance(

            value,

            dict,

        )

        and value.get("id")

    ]



    labels = {}



    try:

        labels = (

            _resolve_entity_labels(

                hq_ids

                + industry_ids

            )

        )

    except Exception:

        labels = {}



    public_lines = [

        line

        for line in (

            company.description

            or ""

        ).splitlines()

        if not line.startswith(

            "Imported public "

        )

    ]



    if websites:

        public_lines.append(

            "Imported public website: "

            f"{websites[0]}"

        )



    if hq_ids:

        public_lines.append(

            "Imported public headquarters: "

            + ", ".join(

                labels.get(

                    item,

                    item,

                )

                for item

                in hq_ids

            )

        )



    if industry_ids:

        public_lines.append(

            "Imported public industries: "

            + ", ".join(

                labels.get(

                    item,

                    item,

                )

                for item

                in industry_ids

            )

        )



    if inception_values:

        first = inception_values[0]



        if isinstance(

            first,

            dict,

        ):

            time_value = (

                first.get("time")

            )



            if time_value:

                public_lines.append(

                    "Imported public inception: "

                    f"{time_value}"

                )



    company.description = (

        "\n".join(public_lines)

    )



    db.commit()



    return {

        "status": "ok",

        "company_id":

            company.id,

        "message": (

            "Public company profile data "

            "was downloaded and stored. "

            "Exact suppliers, internal "

            "routes, route costs, inventory, "

            "daily demand and safety stock "

            "are usually private, so those "

            "values were not invented. "

            "Use Excel/manual entry or "

            "Create Demo Network for them."

        ),

        "source":

            f"https://www.wikidata.org/wiki/{entity_id}",

    }





def _company_country(

    company: Company,

) -> str:

    description = (

        company.description

        or ""

    )



    match = re.search(

        r"Country:\s*([^\n]+)",

        description,

    )



    if match:

        return (

            match.group(1)

            .strip()

        )



    return "Unknown"


def _company_demo_factor(
    company: Company,
    item_key: str,
    minimum: float,
    maximum: float,
) -> float:
    digest = hashlib.sha256(
        f"{company.id}:{company.name.casefold()}:{item_key}".encode("utf-8")
    ).digest()
    fraction = int.from_bytes(digest[:4], "big") / (2**32 - 1)
    return minimum + fraction * (maximum - minimum)


def _demo_node_city(company: Company, node: SupplyNode, index: int) -> str:
    country = _company_country(company)
    locations = COUNTRY_DEMO_NETWORK_CITIES.get(
        country.casefold(),
        (country or "Regional",) * 7,
    )
    standard_node_index = {
        "SUP_01": 0,
        "PLANT_01": 1,
        "PORT_A": 2,
        "PORT_B": 3,
        "DC_01": 4,
        "WH_01": 5,
        "CUST_01": 6,
    }
    city_index = standard_node_index.get(node.code, index % len(locations))
    return f"Example: {locations[city_index]}"


STARTER_DEMO_NODE_LOCATIONS = {
    "SH_SUP": ("Example: Shanghai", "China"),
    "SH_PORT": ("Example: Shanghai", "China"),
    "ROT": ("Example: Rotterdam", "Netherlands"),
    "ANT": ("Example: Antwerp", "Belgium"),
    "HAM": ("Example: Hamburg", "Germany"),
    "FRA_WH": ("Example: Frankfurt", "Germany"),
    "BER": ("Example: Berlin", "Germany"),
}
STARTER_DEMO_ROUTE_CODES = {
    ("SH_SUP", "SH_PORT"),
    ("SH_PORT", "ROT"),
    ("ROT", "FRA_WH"),
    ("SH_PORT", "ANT"),
    ("ANT", "FRA_WH"),
    ("SH_PORT", "HAM"),
    ("HAM", "FRA_WH"),
    ("FRA_WH", "BER"),
}
STARTER_DEMO_PROVENANCE = (
    "Synthetic demo data bundled with SupplyChain Twin AI; "
    "not verified company operating data."
)


def _is_starter_demo_network(
    company: Company,
    nodes: list[SupplyNode],
    routes: list[SupplyRoute],
) -> bool:
    if (
        company.name != "Demo Global Supply Chain"
        or company.description
        != "Synthetic demo workspace shipped with SupplyChain Twin AI."
        or len(nodes) != len(STARTER_DEMO_NODE_LOCATIONS)
        or len(routes) != len(STARTER_DEMO_ROUTE_CODES)
    ):
        return False

    nodes_by_code = {node.code: node for node in nodes}
    if set(nodes_by_code) != set(STARTER_DEMO_NODE_LOCATIONS):
        return False

    route_codes = {
        (
            next(
                (
                    node.code
                    for node in nodes
                    if node.id == route.source_node_id
                ),
                "",
            ),
            next(
                (
                    node.code
                    for node in nodes
                    if node.id == route.destination_node_id
                ),
                "",
            ),
        )
        for route in routes
    }
    return route_codes == STARTER_DEMO_ROUTE_CODES


@router.post(
    "/personalize-demo-network/{company_id}",
    dependencies=[Depends(require_company_access)],
)
def personalize_demo_network(
    company_id: int,
    db: Session = Depends(get_db),
):
    company = db.get(Company, company_id)
    if company is None:
        raise HTTPException(
            status_code=404,
            detail="Company workspace not found.",
        )

    nodes = (
        db.query(SupplyNode)
        .filter(
            SupplyNode.company_id == company_id,
            SupplyNode.active.is_(True),
        )
        .all()
    )
    routes = (
        db.query(SupplyRoute)
        .filter(
            SupplyRoute.company_id == company_id,
            SupplyRoute.active.is_(True),
        )
        .all()
    )

    if not nodes or not routes:
        raise HTTPException(
            status_code=409,
            detail="A complete demo network is required before its values can be varied.",
        )

    if _is_starter_demo_network(company, nodes, routes):
        for node in nodes:
            node.city, node.country = STARTER_DEMO_NODE_LOCATIONS[node.code]
            node.data_source_type = "DEMO / SYNTHETIC / COMPANY VARIANT"
            node.source_url = STARTER_DEMO_PROVENANCE
        for route in routes:
            route.data_source_type = "DEMO / SYNTHETIC / COMPANY VARIANT"
            route.source_reference = STARTER_DEMO_PROVENANCE
        db.commit()
        return {
            "status": "starter_demo_profiled",
            "message": (
                "Example city/country labels and synthetic-data provenance "
                "were added to the bundled demo network. These values are "
                "illustrative, not verified company operating data."
            ),
        }

    nodes_are_synthetic = all(
        "DEMO" in (record.data_source_type or "").upper()
        and "SYNTHETIC" in (record.data_source_type or "").upper()
        for record in nodes
    )
    if not nodes_are_synthetic:
        return {
            "status": "skipped",
            "message": "This workspace contains user-provided or imported node locations, so they were left unchanged.",
        }

    locations_updated = False
    for index, node in enumerate(nodes):
        if not node.city or not node.city.strip():
            node.city = _demo_node_city(company, node, index)
            locations_updated = True

    records = nodes + routes
    all_records_are_synthetic = all(
        "DEMO" in (record.data_source_type or "").upper()
        and "SYNTHETIC" in (record.data_source_type or "").upper()
        for record in records
    )
    if not all_records_are_synthetic:
        if locations_updated:
            db.commit()
            return {
                "status": "locations_updated",
                "message": (
                    "Example city labels were added for synthetic nodes only. "
                    "They are illustrative, not verified facility locations."
                ),
            }
        return {
            "status": "skipped",
            "message": "Imported or user-provided route data was left unchanged.",
        }

    if all(
        "COMPANY VARIANT" in (record.data_source_type or "").upper()
        for record in records
    ):
        if locations_updated:
            db.commit()
            return {
                "status": "locations_updated",
                "message": (
                    "Example city labels were added for this synthetic network. "
                    "They are illustrative only, not verified facility locations."
                ),
            }
        return {
            "status": "unchanged",
            "message": "This synthetic demo network already has company-specific example values.",
        }

    source_type = "DEMO / SYNTHETIC / COMPANY VARIANT"
    for node in nodes:
        node_key = node.code
        node.capacity = max(
            1,
            (node.capacity or 1)
            * _company_demo_factor(company, f"{node_key}:capacity", 0.75, 1.25),
        )
        node.current_inventory *= _company_demo_factor(
            company, f"{node_key}:inventory", 0.65, 1.35
        )
        node.daily_demand *= _company_demo_factor(
            company, f"{node_key}:demand", 0.7, 1.3
        )
        node.safety_stock *= _company_demo_factor(
            company, f"{node_key}:safety-stock", 0.7, 1.3
        )
        node.data_source_type = source_type

    for route in routes:
        route_key = f"{route.source_node_id}->{route.destination_node_id}"
        route.base_cost *= _company_demo_factor(
            company, f"{route_key}:cost", 0.75, 1.3
        )
        route.lead_time_days *= _company_demo_factor(
            company, f"{route_key}:lead-time", 0.8, 1.25
        )
        route.capacity *= _company_demo_factor(
            company, f"{route_key}:capacity", 0.75, 1.25
        )
        route.data_source_type = source_type

    db.commit()
    return {
        "status": "profiled",
        "message": (
            "Synthetic demo values were varied deterministically for this workspace. "
            "They remain examples and are not verified company operating data."
        ),
    }


@router.post(

    "/create-demo-network/{company_id}",
    dependencies=[Depends(require_company_access)],

)

def create_demo_network(

    company_id: int,

    db: Session = Depends(

        get_db

    ),

):

    company = db.get(

        Company,

        company_id,

    )



    if company is None:

        raise HTTPException(

            status_code=404,

            detail=(

                "Company workspace "

                "not found."

            ),

        )



    existing_nodes = (

        db.query(SupplyNode)

        .filter(

            SupplyNode.company_id

            == company_id,

            SupplyNode.active.is_(

                True

            ),

        )

        .count()

    )



    if existing_nodes > 0:

        raise HTTPException(

            status_code=409,

            detail=(

                "This company already has "

                "network nodes. Demo generation "

                "was stopped to avoid mixing "

                "synthetic data with existing data."

            ),

        )



    country = _company_country(

        company

    )



    if country.lower() == "india":

        node_specs = [

            (

                "SUP_01",

                "Primary Components Supplier",

                "SUPPLIER",

                "Pune",

                900,

                0,

                0,

                0,

            ),

            (

                "PLANT_01",

                "Main Manufacturing Plant",

                "PLANT",

                "Pune",

                800,

                0,

                0,

                0,

            ),

            (

                "PORT_A",

                "Primary Gateway Port",

                "PORT",

                "Navi Mumbai",

                700,

                0,

                0,

                0,

            ),

            (

                "PORT_B",

                "Alternate Gateway Port",

                "PORT",

                "Chennai",

                600,

                0,

                0,

                0,

            ),

            (
                "DC_01",
                "Regional Distribution Centre",
                "DISTRIBUTION_CENTER",
                "Hyderabad",
                750,
                6000,

                450,

                1500,

            ),

            (

                "WH_01",

                "Regional Warehouse",

                "WAREHOUSE",

                "Vijayawada",

                550,

                3200,

                300,

                900,

            ),

            (

                "CUST_01",

                "Primary Customer Market",

                "CUSTOMER",

                "Vijayawada",

                500,

                1200,

                350,

                500,

            ),

        ]

    else:

        node_specs = [

            (

                "SUP_01",

                "Primary Supplier",

                "SUPPLIER",

                None,

                900,

                0,

                0,

                0,

            ),

            (

                "PLANT_01",

                "Main Plant",

                "PLANT",

                None,

                800,

                0,

                0,

                0,

            ),

            (

                "PORT_A",

                "Primary Gateway",

                "PORT",

                None,

                700,

                0,

                0,

                0,

            ),

            (

                "PORT_B",

                "Alternate Gateway",

                "PORT",

                None,

                600,

                0,

                0,

                0,

            ),

            (

                "DC_01",

                "Distribution Centre",

                "DISTRIBUTION_CENTER",

                None,

                750,

                6000,

                450,

                1500,

            ),

            (

                "WH_01",

                "Regional Warehouse",

                "WAREHOUSE",

                None,

                550,

                3200,

                300,

                900,

            ),

            (

                "CUST_01",

                "Customer Market",

                "CUSTOMER",

                None,

                500,

                1200,

                350,

                500,

            ),

        ]



    nodes = {}



    for (

        code,

        name,

        node_type,

        city,

        capacity,

        inventory,

        daily_demand,

        safety_stock,

    ) in node_specs:

        node = SupplyNode(

            company_id=company_id,

            code=code,

            name=name,

            node_type=node_type,

            country=(

                None

                if country

                == "Unknown"

                else country

            ),

            city=city,

            latitude=None,

            longitude=None,

            capacity=capacity,

            current_inventory=inventory,

            daily_demand=daily_demand,

            safety_stock=safety_stock,

            data_source_type="DEMO / SYNTHETIC",

            active=True,

        )

        db.add(node)

        db.flush()



        nodes[code] = node



    route_specs = [

        (

            "SUP_01",

            "PLANT_01",

            "road",

            1200,

            1.0,

            800,

            1.0,

        ),

        (

            "PLANT_01",

            "PORT_A",

            "road",

            2200,

            1.5,

            700,

            1.0,

        ),

        (

            "PORT_A",

            "DC_01",

            "rail",

            5200,

            3.0,

            600,

            0.4,

        ),

        (

            "PLANT_01",

            "PORT_B",

            "rail",

            4800,

            3.0,

            600,

            0.4,

        ),

        (

            "PORT_B",

            "DC_01",

            "road",

            6500,

            2.5,

            500,

            1.0,

        ),

        (

            "DC_01",

            "WH_01",

            "road",

            2600,

            1.0,

            500,

            1.0,

        ),

        (

            "WH_01",

            "CUST_01",

            "road",

            900,

            0.5,

            450,

            1.0,

        ),

        (

            "DC_01",

            "CUST_01",

            "road",

            3300,

            1.5,

            400,

            1.0,

        ),

    ]



    for (

        source_code,

        destination_code,

        mode,

        cost,

        lead_time,

        capacity,

        fuel_dependency,

    ) in route_specs:

        db.add(

            SupplyRoute(

                company_id=company_id,

                source_node_id=(

                    nodes[

                        source_code

                    ].id

                ),

                destination_node_id=(

                    nodes[

                        destination_code

                    ].id

                ),

                transport_mode=mode,

                base_cost=cost,

                lead_time_days=lead_time,

                capacity=capacity,

                fuel_dependency=fuel_dependency,

                data_source_type="DEMO / SYNTHETIC",

                active=True,

            )

        )

    company.is_demo = True

    db.commit()



    return {

        "status": "ok",

        "nodes_created":

            len(node_specs),

        "routes_created":

            len(route_specs),

        "message": (

            "Synthetic demo network created. "

            "These operational values are estimated/demo data and are intended for app testing and visualization only."

        ),

    }





def _normalize_header(

    value,

) -> str:

    return (

        str(value or "")

        .strip()

        .lower()

        .replace(" ", "_")

        .replace("-", "_")

    )





def _sheet_records(

    sheet,

) -> list[dict]:

    rows = list(

        sheet.iter_rows(

            values_only=True

        )

    )



    if not rows:

        return []



    headers = [

        _normalize_header(

            value

        )

        for value in rows[0]

    ]



    records = []



    for row_number, row in enumerate(rows[1:], start=2):

        if not any(

            value is not None

            and str(value).strip()

            for value in row

        ):

            continue



        record = {}
        record["__row__"] = row_number



        for index, header in enumerate(

            headers

        ):

            if not header:

                continue



            record[header] = (

                row[index]

                if index < len(row)

                else None

            )



        records.append(record)



    return records


def _read_excel_network(raw: bytes) -> tuple[list[dict], list[dict]]:
    if not raw:
        raise HTTPException(
            status_code=400,
            detail="The uploaded Excel file is empty.",
        )

    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise HTTPException(
            status_code=500,
            detail="Excel import requires openpyxl. Install it in the backend environment.",
        ) from error

    try:
        workbook = load_workbook(
            filename=BytesIO(raw),
            data_only=True,
            read_only=True,
        )
    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail=f"Unable to read Excel file: {error}",
        ) from error

    sheet_lookup = {
        name.lower(): workbook[name]
        for name in workbook.sheetnames
    }
    nodes_sheet = sheet_lookup.get("nodes")
    routes_sheet = sheet_lookup.get("routes")
    if nodes_sheet is None or routes_sheet is None:
        missing = [
            name
            for name, sheet in (("Nodes", nodes_sheet), ("Routes", routes_sheet))
            if sheet is None
        ]
        raise HTTPException(
            status_code=400,
            detail=(
                "Excel workbook must contain both Nodes and Routes sheets. "
                f"Missing: {', '.join(missing)}."
            ),
        )

    return _sheet_records(nodes_sheet), _sheet_records(routes_sheet)


def _validate_excel_network(
    node_records: list[dict],
    route_records: list[dict],
    existing_node_codes: set[str],
) -> list[str]:
    errors: list[str] = []
    allowed_node_types = {
        "SUPPLIER", "PLANT", "PORT", "WAREHOUSE",
        "DISTRIBUTION_CENTER", "CUSTOMER", "OTHER",
    }
    optional_node_numbers = (
        "latitude", "longitude", "capacity",
        "current_inventory", "daily_demand", "safety_stock",
    )
    seen_codes: set[str] = set()

    if not node_records:
        errors.append("Nodes sheet has no data rows.")
    if not route_records:
        errors.append("Routes sheet has no data rows.")

    for index, item in enumerate(node_records, start=2):
        row = item.get("__row__", index)
        code = str(item.get("code") or "").strip().upper()
        name = str(item.get("name") or "").strip()
        node_type = str(item.get("node_type") or "").strip().upper()
        item["code"] = code
        item["name"] = name
        item["node_type"] = node_type
        if not code:
            errors.append(f"Nodes row {row}: code is required.")
        elif code in seen_codes:
            errors.append(f"Nodes row {row}: duplicate node code '{code}'.")
        seen_codes.add(code)
        if not name:
            errors.append(f"Nodes row {row}: name is required.")
        if node_type not in allowed_node_types:
            errors.append(
                f"Nodes row {row}: node_type must be one of "
                f"{', '.join(sorted(allowed_node_types))}."
            )

        for field in optional_node_numbers:
            value = item.get(field)
            if field == "capacity" and (
                value is None or (isinstance(value, str) and not value.strip())
            ):
                errors.append(f"Nodes row {row}: capacity is required.")
                continue
            if value is None or (isinstance(value, str) and not value.strip()):
                item[field] = 0.0 if field in {
                    "current_inventory", "daily_demand", "safety_stock"
                } else None
                continue
            try:
                parsed = float(value)
            except (TypeError, ValueError):
                errors.append(f"Nodes row {row}: {field} must be a number.")
                continue
            if not math.isfinite(parsed):
                errors.append(f"Nodes row {row}: {field} must be a finite number.")
            elif field == "latitude" and not -90 <= parsed <= 90:
                errors.append(f"Nodes row {row}: latitude must be between -90 and 90.")
            elif field == "longitude" and not -180 <= parsed <= 180:
                errors.append(f"Nodes row {row}: longitude must be between -180 and 180.")
            elif field in {"current_inventory", "daily_demand", "safety_stock"} and parsed < 0:
                errors.append(f"Nodes row {row}: {field} cannot be negative.")
            elif field == "capacity" and parsed <= 0:
                errors.append(f"Nodes row {row}: capacity must be greater than zero.")
            item[field] = parsed

    available_codes = existing_node_codes | seen_codes
    seen_routes: set[tuple[str, str]] = set()
    required_route_numbers = {
        "base_cost": ("base_cost", "cost"),
        "lead_time_days": ("lead_time_days", "lead_time"),
        "capacity": ("capacity",),
    }
    for index, item in enumerate(route_records, start=2):
        row = item.get("__row__", index)
        source = str(item.get("source_code") or "").strip().upper()
        destination = str(item.get("destination_code") or "").strip().upper()
        item["source_code"] = source
        item["destination_code"] = destination
        if not source or not destination:
            errors.append(
                f"Routes row {row}: source_code and destination_code are required."
            )
        elif source == destination:
            errors.append(f"Routes row {row}: source and destination must differ.")
        else:
            pair = (source, destination)
            if pair in seen_routes:
                errors.append(
                    f"Routes row {row}: duplicate route '{source} -> {destination}'."
                )
            seen_routes.add(pair)
        for endpoint in (source, destination):
            if endpoint and endpoint not in available_codes:
                errors.append(
                    f"Routes row {row}: node code '{endpoint}' is not in the Nodes sheet or this workspace."
                )

        for normalized, aliases in required_route_numbers.items():
            value = next(
                (item.get(alias) for alias in aliases if item.get(alias) is not None),
                None,
            )
            try:
                parsed = float(value)
            except (TypeError, ValueError):
                errors.append(f"Routes row {row}: {normalized} must be a number.")
                continue
            if not math.isfinite(parsed):
                errors.append(f"Routes row {row}: {normalized} must be a finite number.")
            elif parsed <= 0:
                errors.append(f"Routes row {row}: {normalized} must be greater than zero.")
            item[normalized] = parsed

        dependency = item.get("fuel_dependency")
        if dependency is None or (isinstance(dependency, str) and not dependency.strip()):
            item["fuel_dependency"] = 1.0
        else:
            try:
                parsed_dependency = float(dependency)
            except (TypeError, ValueError):
                errors.append(f"Routes row {row}: fuel_dependency must be a number between 0 and 1.")
            else:
                if not math.isfinite(parsed_dependency) or not 0 <= parsed_dependency <= 1:
                    errors.append(f"Routes row {row}: fuel_dependency must be between 0 and 1.")
                item["fuel_dependency"] = parsed_dependency

    if not errors and len(available_codes) < 2:
        errors.append("At least two distinct nodes are required to define routes.")
    return errors


@router.get(
    "/excel-template"
)
def download_excel_template():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise HTTPException(
            status_code=500,
            detail="Excel template generation requires openpyxl.",
        ) from error

    workbook = Workbook()
    nodes = workbook.active
    nodes.title = "Nodes"
    nodes.append([
        "code", "name", "node_type", "country", "city",
        "latitude", "longitude", "capacity", "current_inventory",
        "daily_demand", "safety_stock",
    ])
    routes = workbook.create_sheet("Routes")
    routes.append([
        "source_code", "destination_code", "transport_mode",
        "base_cost", "lead_time_days", "capacity", "fuel_dependency",
    ])
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Workbook field", "Requirement"])
    instructions.append(["Node code, name, node_type", "Required on every Nodes row; codes must be unique."])
    instructions.append(["Node type", "SUPPLIER, PLANT, PORT, WAREHOUSE, DISTRIBUTION_CENTER, CUSTOMER, or OTHER."])
    instructions.append(["Node capacity", "Required and must be greater than zero. Inventory and demand values cannot be negative."])
    instructions.append(["Route source_code, destination_code", "Required; both codes must exist in this workbook or the selected workspace."])
    instructions.append(["Route base_cost, lead_time_days, capacity", "Required and must be greater than zero."])
    instructions.append(["Route fuel_dependency", "Optional from 0 to 1; defaults to 1."])

    for sheet in (nodes, routes, instructions):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0E7490")
        for column in sheet.columns:
            width = min(
                max(len(str(cell.value or "")) for cell in column) + 3,
                55,
            )
            sheet.column_dimensions[column[0].column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="supply-chain-template.xlsx"',
        },
    )


@router.post(
    "/preview-excel/{company_id}",
    dependencies=[Depends(require_company_access)],
)
async def preview_excel(
    company_id: int,
    request: Request,
    filename: str = "network.xlsx",
    db: Session = Depends(get_db),
):
    if db.get(Company, company_id) is None:
        raise HTTPException(status_code=404, detail="Company workspace not found.")

    node_records, route_records = _read_excel_network(await request.body())
    existing_codes = {
        node.code.strip().upper()
        for node in db.query(SupplyNode)
        .filter(SupplyNode.company_id == company_id)
        .all()
    }
    errors = _validate_excel_network(
        node_records, route_records, existing_codes
    )
    return {
        "filename": filename,
        "valid": not errors,
        "nodes_count": len(node_records),
        "routes_count": len(route_records),
        "errors": errors,
        "nodes_preview": [
            {
                "code": item.get("code"),
                "name": item.get("name"),
                "node_type": item.get("node_type"),
            }
            for item in node_records[:5]
        ],
        "routes_preview": [
            {
                "source_code": item.get("source_code"),
                "destination_code": item.get("destination_code"),
                "transport_mode": item.get("transport_mode") or "road",
                "base_cost": item.get("base_cost"),
                "lead_time_days": item.get("lead_time_days"),
            }
            for item in route_records[:5]
        ],
    }


@router.post(

    "/import-excel/{company_id}",
    dependencies=[Depends(require_company_access)],

)

async def import_excel(

    company_id: int,

    request: Request,

    filename: str = "network.xlsx",

    db: Session = Depends(

        get_db

    ),

):

    company = db.get(

        Company,

        company_id,

    )



    if company is None:

        raise HTTPException(

            status_code=404,

            detail=(

                "Company workspace "

                "not found."

            ),

        )



    try:

        from openpyxl import (

            load_workbook,

        )

    except ImportError:

        raise HTTPException(

            status_code=500,

            detail=(

                "Excel import requires "

                "openpyxl. Run: "

                ".\\\\.venv\\\Scripts\\\python.exe "

                "-m pip install openpyxl"

            ),

        )



    raw = await request.body()



    if not raw:

        raise HTTPException(

            status_code=400,

            detail=(

                "The uploaded Excel file "

                "is empty."

            ),

        )



    try:

        workbook = load_workbook(

            filename=BytesIO(raw),

            data_only=True,

        )

    except Exception as error:

        raise HTTPException(

            status_code=400,

            detail=(

                "Unable to read Excel "

                f"file: {error}"

            ),

        )



    sheet_lookup = {

        name.lower():

            workbook[name]

        for name

        in workbook.sheetnames

    }



    nodes_sheet = (

        sheet_lookup.get(

            "nodes"

        )

    )



    routes_sheet = (

        sheet_lookup.get(

            "routes"

        )

    )



    if nodes_sheet is None:

        raise HTTPException(

            status_code=400,

            detail=(

                "Excel must contain a "

                "sheet named Nodes."

            ),

        )



    node_records = (

        _sheet_records(

            nodes_sheet

        )

    )



    route_records = (

        _sheet_records(

            routes_sheet

        )

        if routes_sheet

        is not None

        else []
    )

    existing_node_codes = {
        node.code.strip().upper()
        for node in db.query(SupplyNode)
        .filter(SupplyNode.company_id == company_id)
        .all()
    }
    validation_errors = _validate_excel_network(
        node_records,
        route_records,
        existing_node_codes,
    )
    if validation_errors:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Excel workbook has validation errors. No data was imported.",
                "errors": validation_errors,
            },
        )



    nodes_by_code = {

        node.code:

            node

        for node in (

            db.query(SupplyNode)

            .filter(

                SupplyNode.company_id

                == company_id

            )

            .all()

        )

    }



    nodes_imported = 0



    for item in node_records:

        code = str(

            item.get(

                "code"

            )

            or ""

        ).strip().upper()



        name = str(

            item.get(

                "name"

            )

            or ""

        ).strip()



        node_type = str(

            item.get(

                "node_type"

            )

            or ""

        ).strip().upper()



        if not (

            code

            and name

            and node_type

        ):

            continue



        node = (

            nodes_by_code.get(

                code

            )

        )



        values = {

            "name":

                name,

            "node_type":

                node_type,

            "country":

                item.get(

                    "country"

                ),

            "city":

                item.get(

                    "city"

                ),

            "latitude":

                item.get(

                    "latitude"

                ),

            "longitude":

                item.get(

                    "longitude"

                ),

            "capacity":

                item.get(

                    "capacity"

                ),

            "current_inventory":

                item.get(

                    "current_inventory"

                )

                or 0,

            "daily_demand":

                item.get(

                    "daily_demand"

                )

                or 0,

            "safety_stock":

                item.get(

                    "safety_stock"

                )

                or 0,

            "active":

                True,

            "data_source_type":

                "IMPORTED FROM EXCEL",

        }



        if node is None:

            node = SupplyNode(

                company_id=company_id,

                code=code,

                **values,

            )



            db.add(node)

            db.flush()



            nodes_by_code[

                code

            ] = node

        else:

            for key, value in (

                values.items()

            ):

                setattr(

                    node,

                    key,

                    value,

                )



        nodes_imported += 1



    routes_imported = 0



    for item in route_records:

        source_code = str(

            item.get(

                "source_code"

            )

            or ""

        ).strip().upper()



        destination_code = str(

            item.get(

                "destination_code"

            )

            or ""

        ).strip().upper()



        source = nodes_by_code.get(

            source_code

        )



        destination = (

            nodes_by_code.get(

                destination_code

            )

        )



        if (

            source is None

            or destination is None

        ):

            continue



        existing = (

            db.query(SupplyRoute)

            .filter(

                SupplyRoute.company_id

                    == company_id,

                SupplyRoute.source_node_id

                    == source.id,

                SupplyRoute.destination_node_id

                    == destination.id,

            )

            .first()

        )



        values = {

            "transport_mode":

                str(

                    item.get(

                        "transport_mode"

                    )

                    or "road"

                ).strip().lower(),

            "base_cost":
                item["base_cost"],

            "lead_time_days":
                item["lead_time_days"],

            "capacity":
                item["capacity"],

            "fuel_dependency":

                item.get(

                    "fuel_dependency"

                )

                if item.get(

                    "fuel_dependency"

                )

                is not None

                else 1,

            "active":

                True,

            "data_source_type":

                "IMPORTED FROM EXCEL",

        }



        if existing is None:

            db.add(

                SupplyRoute(

                    company_id=company_id,

                    source_node_id=source.id,

                    destination_node_id=(

                        destination.id

                    ),

                    **values,

                )

            )

        else:

            for key, value in (

                values.items()

            ):

                setattr(

                    existing,

                    key,

                    value,

                )



        routes_imported += 1



    db.commit()



    return {

        "status": "ok",

        "filename":

            filename,

        "nodes_imported":

            nodes_imported,

        "routes_imported":

            routes_imported,

    }
