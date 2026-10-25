import csv
import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.auth import (
    get_current_user,
    log_workspace_activity,
    require_company_read_access,
    require_company_write_access,
)
from app.db.session import get_db
from app.db_models.entities import (
    SupplierPerformanceRecord,
    SupplierRiskThreshold,
    SupplyNode,
    User,
)
from app.schemas.network import SupplierPerformanceCreate, SupplierPerformanceRead
from app.schemas.network import SupplierRiskThresholdRead
from app.supplier_risk import build_supplier_risk_alerts, notify_supplier_risk_alerts

REPORT_DATE_CONSTRAINT = "uq_supplier_performance_reporting_date"

router = APIRouter(
    prefix="/companies/{company_id}/supplier-performance",
    tags=["supplier performance"],
)


def require_company_access(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_company_read_access(company_id, current_user, db)


@router.get(
    "",
    response_model=list[SupplierPerformanceRead],
    dependencies=[Depends(require_company_access)],
)
def list_supplier_performance(
    company_id: int,
    supplier_node_id: int | None = Query(default=None, gt=0),
    start_date: date | None = None,
    end_date: date | None = None,
    limit: int = Query(default=500, ge=1, le=5000),
    db: Session = Depends(get_db),
):
    if start_date is not None and end_date is not None and start_date > end_date:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Start date must be on or before end date.",
        )

    query = (
        db.query(SupplierPerformanceRecord)
        .join(SupplyNode, SupplierPerformanceRecord.supplier_node_id == SupplyNode.id)
        .filter(
            SupplyNode.company_id == company_id,
            SupplyNode.node_type.ilike("supplier"),
        )
    )
    if supplier_node_id is not None:
        query = query.filter(
            SupplierPerformanceRecord.supplier_node_id == supplier_node_id
        )
    if start_date is not None:
        query = query.filter(SupplierPerformanceRecord.reporting_date >= start_date)
    if end_date is not None:
        query = query.filter(SupplierPerformanceRecord.reporting_date <= end_date)

    return (
        query.order_by(
            SupplierPerformanceRecord.reporting_date.desc(),
            SupplierPerformanceRecord.id.desc(),
        )
        .limit(limit)
        .all()
    )


@router.post(
    "",
    response_model=SupplierPerformanceRead,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_company_write_access)],
)
def create_supplier_performance(
    company_id: int,
    payload: SupplierPerformanceCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    supplier = db.get(SupplyNode, payload.supplier_node_id)
    if (
        supplier is None
        or supplier.company_id != company_id
        or supplier.node_type.casefold() != "supplier"
    ):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Supplier location not found in this workspace.",
        )

    existing_record = (
        db.query(SupplierPerformanceRecord.id)
        .filter_by(
            supplier_node_id=supplier.id,
            reporting_date=payload.reporting_date,
        )
        .first()
    )
    if existing_record is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A performance record already exists for this supplier and reporting date.",
        )

    record = SupplierPerformanceRecord(
        **payload.model_dump(),
        recorded_by_id=current_user.id,
    )
    db.add(record)
    log_workspace_activity(
        db,
        company_id,
        current_user,
        "supplier_performance_recorded",
        f"Recorded supplier performance for {supplier.name} on {payload.reporting_date.isoformat()}.",
    )
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        constraint_name = getattr(
            getattr(error.orig, "diag", None),
            "constraint_name",
            None,
        )
        sqlite_duplicate = (
            "UNIQUE constraint failed: "
            "supplier_performance_records.supplier_node_id, "
            "supplier_performance_records.reporting_date"
        )
        if constraint_name != REPORT_DATE_CONSTRAINT and sqlite_duplicate not in str(
            error.orig
        ):
            raise
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="A performance record already exists for this supplier and reporting date.",
        ) from error
    db.refresh(record)
    thresholds = (
        db.query(SupplierRiskThreshold)
        .filter_by(company_id=company_id)
        .first()
    )
    if thresholds is not None:
        threshold_read = SupplierRiskThresholdRead.model_validate(thresholds)
        alerts = [
            alert
            for alert in build_supplier_risk_alerts(db, company_id, threshold_read)
            if alert.performance_record_id == record.id
        ]
        notify_supplier_risk_alerts(
            db,
            company_id,
            current_user.id,
            alerts,
            threshold_read.updated_at,
        )
        db.commit()
    return record


def _normalize_import_key(value: str | None) -> str:
    if value is None:
        return ""
    return "".join(ch.lower() if ch.isalnum() else "_" for ch in value.strip()).strip("_")


def _coerce_float(value: object, field_name: str) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text in {"", "-", "--", "n/a", "na", "null", "none", "nan", "inf", "-inf", "—"}:
        return None
    normalized = text.replace(",", "")
    if normalized.lower() in {"n/a", "na", "null", "none"}:
        return None
    try:
        parsed = float(normalized)
    except ValueError as exc:  # pragma: no cover - defensive validation
        raise ValueError(f"{field_name} must be numeric.") from exc
    return parsed


def _read_supplier_import_rows(file_bytes: bytes, filename: str) -> list[dict[str, object]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows: list[dict[str, object]] = []
        for row in reader:
            if row is None:
                continue
            if all(str(value or "").strip() == "" for value in row.values()):
                continue
            rows.append({str(key): value for key, value in row.items()})
        return rows

    if name.endswith((".xlsx", ".xlsm", ".xls")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - runtime dependency guard
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Excel import requires openpyxl. Install it in the backend environment.",
            ) from exc

        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        rows: list[dict[str, object]] = []
        header_row: list[str] | None = None
        for worksheet in workbook.worksheets:
            for sheet_row in worksheet.iter_rows(values_only=True):
                if not sheet_row or all(
                    value is None or str(value).strip() == "" for value in sheet_row
                ):
                    continue
                values = list(sheet_row)
                if header_row is None:
                    header_row = [str(value).strip() for value in values]
                    continue
                row_data: dict[str, object] = {}
                for header, value in zip(header_row, values):
                    if header:
                        row_data[header] = value
                if row_data:
                    rows.append(row_data)
        return rows

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Supplier performance uploads must be CSV or Excel files.",
    )


def _map_import_row(raw_row: dict[str, object]) -> dict[str, object]:
    mapped: dict[str, object] = {}
    for key, value in raw_row.items():
        normalized = _normalize_import_key(str(key))
        if normalized:
            mapped[normalized] = value
    return mapped


def _first_non_empty(*values: object) -> str | None:
    for value in values:
        if value is None:
            continue
        text = str(value).strip()
        if text not in {"", "-", "--", "null", "none", "n/a", "na", "—"}:
            return text
    return None


def _parse_supplier_performance_row(
    company_id: int,
    db: Session,
    raw_row: dict[str, object],
) -> tuple[dict[str, object], list[str]]:
    row = _map_import_row(raw_row)
    errors: list[str] = []

    direct_supplier_id = row.get("supplier_node_id")
    if direct_supplier_id is not None:
        try:
            supplier_id = int(direct_supplier_id)
        except (TypeError, ValueError):
            errors.append("Supplier identifier must be an integer.")
            supplier_id = None
        else:
            supplier = db.get(SupplyNode, supplier_id)
            if (
                supplier is None
                or supplier.company_id != company_id
                or supplier.node_type.casefold() != "supplier"
            ):
                errors.append("Supplier was not found in this workspace.")
                supplier_id = None
    else:
        supplier_code = _first_non_empty(
            row.get("supplier_code"),
            row.get("supplier"),
            row.get("supplier_name"),
            row.get("supplierid"),
        )
        if supplier_code is None:
            errors.append("Missing supplier code or supplier name.")
            supplier_id = None
        else:
            suppliers = (
                db.query(SupplyNode)
                .filter(
                    SupplyNode.company_id == company_id,
                    SupplyNode.node_type.ilike("supplier"),
                )
                .all()
            )
            supplier = next(
                (
                    candidate
                    for candidate in suppliers
                    if candidate.code.strip().casefold() == supplier_code.casefold()
                    or candidate.name.strip().casefold() == supplier_code.casefold()
                ),
                None,
            )
            if supplier is None:
                errors.append(f"Supplier '{supplier_code}' was not found in this workspace.")
                supplier_id = None
            else:
                supplier_id = supplier.id

    reporting_date_text = _first_non_empty(
        row.get("reporting_date"),
        row.get("date"),
        row.get("report_date"),
    )
    if reporting_date_text is None:
        errors.append("Missing reporting date.")
        reporting_date = None
    else:
        try:
            reporting_date = date.fromisoformat(reporting_date_text)
        except ValueError:
            errors.append("Reporting date must be a valid ISO date (YYYY-MM-DD).")
            reporting_date = None

    metric_keys = {
        "on_time_delivery_rate": (
            "on_time_delivery_rate",
            "on_time_delivery",
            "on_time_delivery_percent",
            "on_timedelivery",
        ),
        "average_lead_time_days": (
            "average_lead_time_days",
            "average_lead_time",
            "lead_time_days",
            "lead_time",
            "averageleadtimedays",
        ),
        "defect_rate": ("defect_rate", "defect", "defect_rate_percent", "defectpercent"),
        "capacity_utilization": (
            "capacity_utilization",
            "capacity",
            "capacity_utilization_percent",
            "capacitypercent",
        ),
    }

    values: dict[str, float | None] = {}
    for metric_name, aliases in metric_keys.items():
        raw_value = None
        for alias in aliases:
            if alias in row:
                raw_value = row[alias]
                break
        if raw_value is None:
            values[metric_name] = None
            continue
        try:
            values[metric_name] = _coerce_float(raw_value, metric_name.replace("_", " "))
        except ValueError as exc:
            errors.append(str(exc))
            values[metric_name] = None

    if all(value is None for value in values.values()):
        errors.append("Enter at least one supplier performance metric.")

    source_reference = _first_non_empty(
        row.get("source_reference"),
        row.get("source"),
        row.get("reference"),
    )
    notes = _first_non_empty(
        row.get("notes"),
        row.get("note"),
    )

    if supplier_id is None or reporting_date is None:
        record: dict[str, object] = {
            "supplier_node_id": supplier_id,
            "reporting_date": reporting_date.isoformat() if reporting_date else None,
            "source_reference": source_reference,
            "notes": notes,
        }
        for metric_name in values:
            record[metric_name] = values[metric_name]
        return record, errors

    record = {
        "supplier_node_id": supplier_id,
        "reporting_date": reporting_date.isoformat(),
        "on_time_delivery_rate": values["on_time_delivery_rate"],
        "average_lead_time_days": values["average_lead_time_days"],
        "defect_rate": values["defect_rate"],
        "capacity_utilization": values["capacity_utilization"],
        "source_reference": source_reference,
        "notes": notes,
    }

    try:
        SupplierPerformanceCreate(**record)
    except ValueError as exc:
        for message in str(exc).split("\n"):
            if message.strip():
                errors.append(message.strip())

    existing = (
        db.query(SupplierPerformanceRecord.id)
        .filter_by(
            supplier_node_id=supplier_id,
            reporting_date=reporting_date,
        )
        .first()
    )
    if existing is not None:
        errors.append("A performance record already exists for this supplier and reporting date.")

    return record, errors


@router.post(
    "/import-preview",
    dependencies=[Depends(require_company_write_access)],
)
async def preview_supplier_performance_import(
    company_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    content = await file.read()
    try:
        raw_rows = _read_supplier_import_rows(content, file.filename or "")
    except HTTPException:
        raise

    valid_records: list[dict[str, object]] = []
    preview_rows: list[dict[str, object]] = []
    for index, raw_row in enumerate(raw_rows, start=1):
        parsed, errors = _parse_supplier_performance_row(company_id, db, raw_row)
        if errors:
            preview_rows.append(
                {
                    "row_number": index,
                    "valid": False,
                    "errors": errors,
                    "data": raw_row,
                }
            )
            continue
        valid_records.append(parsed)
        preview_rows.append(
            {
                "row_number": index,
                "valid": True,
                "errors": [],
                "data": parsed,
            }
        )

    return {
        "filename": file.filename,
        "total_rows": len(raw_rows),
        "valid_rows": len(valid_records),
        "invalid_rows": len(raw_rows) - len(valid_records),
        "valid_records": valid_records,
        "rows": preview_rows,
    }


@router.post(
    "/import",
    dependencies=[Depends(require_company_write_access)],
)
def import_supplier_performance_records(
    company_id: int,
    payload: dict[str, list[dict[str, object]]],
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    rows = payload.get("rows", [])
    if not isinstance(rows, list):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Import payload must contain a 'rows' list.",
        )

    imported_records: list[SupplierPerformanceRecord] = []
    errors: list[dict[str, object]] = []
    seen_duplicates: set[tuple[int, date]] = set()

    for index, raw_row in enumerate(rows, start=1):
        try:
            parsed, parse_errors = _parse_supplier_performance_row(company_id, db, raw_row)
        except Exception as exc:  # pragma: no cover - defensive runtime guard
            errors.append({"row_number": index, "errors": [str(exc)]})
            continue
        if parse_errors:
            errors.append({"row_number": index, "errors": parse_errors})
            continue

        supplier_id = int(parsed["supplier_node_id"])
        reporting_date = date.fromisoformat(str(parsed["reporting_date"]))
        key = (supplier_id, reporting_date)
        if key in seen_duplicates:
            errors.append({
                "row_number": index,
                "errors": ["A performance record already exists for this supplier and reporting date."],
            })
            continue
        seen_duplicates.add(key)

        record = SupplierPerformanceRecord(
            supplier_node_id=supplier_id,
            reporting_date=reporting_date,
            on_time_delivery_rate=parsed.get("on_time_delivery_rate"),
            average_lead_time_days=parsed.get("average_lead_time_days"),
            defect_rate=parsed.get("defect_rate"),
            capacity_utilization=parsed.get("capacity_utilization"),
            source_reference=parsed.get("source_reference"),
            notes=parsed.get("notes"),
            recorded_by_id=current_user.id,
        )
        db.add(record)
        imported_records.append(record)

    if not imported_records:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No valid supplier performance rows were available to import.",
        )

    db.commit()
    for record in imported_records:
        db.refresh(record)

    return {
        "imported_count": len(imported_records),
        "errors": errors,
        "records": [
            {
                "id": record.id,
                "supplier_node_id": record.supplier_node_id,
                "reporting_date": record.reporting_date.isoformat(),
                "on_time_delivery_rate": record.on_time_delivery_rate,
                "average_lead_time_days": record.average_lead_time_days,
                "defect_rate": record.defect_rate,
                "capacity_utilization": record.capacity_utilization,
                "source_reference": record.source_reference,
                "notes": record.notes,
            }
            for record in imported_records
        ],
    }
